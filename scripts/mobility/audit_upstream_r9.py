"""Offline/online upstream drift audit for W1 pinned mobility evidence."""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from scripts.mobility.source_watch_r9 import fetch


ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / "docs/vnext/w1"
RAW = ROOT / "datos_originales"
PINNED = {
    "GTFS": RAW / "movilidad/goierrialdea-3276fcae.zip",
    "HEALTH_REGISTRY": RAW / "centros-salud.xlsx",
    "HEALTH_PAGE": RAW / "movilidad/r5/beasain-official.html",
    "PADI_2026": RAW / "movilidad/r5/padi-2026.pdf",
    "OSM": RAW / "movilidad/beasain-network-r4-public.osm",
}
URLS = {
    "GTFS": "https://opendata.euskadi.eus/transport/moveuskadi/lurraldebus/goierrialdea/gtfs_goierrialdea.zip",
    "HEALTH_REGISTRY": "https://opendata.euskadi.eus/contenidos/ds_localizaciones/centros_salud_en_euskadi/opendata/centros-salud.xlsx",
    "HEALTH_PAGE": "https://www.osakidetza.euskadi.eus/ambulatorio-de-beasain/webosk00-cercon/es/",
    "PADI_2026": "https://www.osakidetza.euskadi.eus/contenidos/informacion/salud_padi/es_def/adjuntos/padi-kontsultak-gipuzkoa.pdf",
    "OSM": "https://api.openstreetmap.org/api/0.6/map?bbox=-2.203,43.041,-2.190,43.052",
}
TABLES = ("agency.txt", "routes.txt", "trips.txt", "stops.txt", "stop_times.txt", "calendar.txt", "calendar_dates.txt")
DRIFT = {"UNCHANGED_BYTES", "CHANGED_BYTES_SEMANTICALLY_EQUIVALENT_USED_FIELDS", "CHANGED_METADATA_ONLY",
         "CHANGED_USED_FIELDS", "SOURCE_UNAVAILABLE", "INDETERMINATE"}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_sha(path: Path, source_id: str) -> str:
    raw = path.read_bytes()
    if source_id == "OSM":
        raw = raw.replace(b"\r\n", b"\n").replace(b"\r", b"\n")
    return hashlib.sha256(raw).hexdigest()


def canonical_rows(archive: zipfile.ZipFile, name: str) -> list[dict[str, str]]:
    if name not in archive.namelist():
        return []
    with archive.open(name) as stream:
        rows = list(csv.DictReader(io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")))
    return sorted(({str(k): str(v) for k, v in row.items()} for row in rows), key=lambda row: tuple(sorted(row.items())))


def stable_hash(value) -> str:
    raw = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode()
    return hashlib.sha256(raw).hexdigest()


def gtfs_signature(path: Path, used_trip_ids: set[str], used_stop_ids: set[str]) -> dict:
    with zipfile.ZipFile(path) as archive:
        tables = {name: canonical_rows(archive, name) for name in TABLES}
    trips = {row.get("trip_id"): row for row in tables["trips.txt"] if row.get("trip_id") in used_trip_ids}
    service_ids = {row.get("service_id") for row in trips.values()}
    signature = {
        "used_trips": trips,
        "used_stops": {row.get("stop_id"): row for row in tables["stops.txt"] if row.get("stop_id") in used_stop_ids},
        "used_stop_times": [row for row in tables["stop_times.txt"] if row.get("trip_id") in used_trip_ids and row.get("stop_id") in used_stop_ids],
        "used_calendar": [row for row in tables["calendar.txt"] if row.get("service_id") in service_ids],
        "used_calendar_dates": [row for row in tables["calendar_dates.txt"] if row.get("service_id") in service_ids and row.get("date") == "20260929"],
        "go01_routes": [row for row in tables["routes.txt"] if row.get("route_short_name", "").upper() == "GO01"],
    }
    return {"whole_sha256": sha(path), "canonical_table_sha256": {name: stable_hash(rows) for name, rows in tables.items()},
            "used_fields_sha256": stable_hash(signature), "used_fields": signature}


def health_record(path: Path) -> dict:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(value or "") for value in next(rows)]
    record = next(dict(zip(header, row)) for row in rows if row[1] == "entityBC631DA5")
    def pick(prefix):
        key = next(key for key in record if key.startswith(prefix))
        return str(record[key] or "").strip()
    return {"entity_id": str(record[header[1]]), "name": pick("Nombre"), "address": pick("Direcci"),
            "latitude": pick("LATWGS84"), "longitude": pick("LONWGS84"), "telephone": pick("Tel")}


def page_signature(path: Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    clean = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", text))).strip().casefold()
    values = {"name": "ambulatorio de beasain", "address": "bernedo enea, 1", "telephone": "943027700"}
    return {key: value in clean for key, value in values.items()}


def osm_signature(path: Path, node_ids: set[str], way_ids: set[str]) -> dict:
    root = ET.fromstring(path.read_bytes())
    def tags(item): return sorted((tag.get("k"), tag.get("v")) for tag in item.findall("tag"))
    nodes = {item.get("id"): {"lat": item.get("lat"), "lon": item.get("lon"), "tags": tags(item)}
             for item in root.findall("node") if item.get("id") in node_ids}
    ways = {item.get("id"): {"nodes": [n.get("ref") for n in item.findall("nd")], "tags": tags(item)}
            for item in root.findall("way") if item.get("id") in way_ids}
    return {"used_nodes": nodes, "used_ways": ways, "used_fields_sha256": stable_hash({"nodes": nodes, "ways": ways})}


def classification(pinned_hash: str, current_hash: str, pinned_semantic, current_semantic, *, metadata_only=False) -> str:
    if pinned_hash == current_hash:
        return "UNCHANGED_BYTES"
    if pinned_semantic == current_semantic:
        return "CHANGED_METADATA_ONLY" if metadata_only else "CHANGED_BYTES_SEMANTICALLY_EQUIVALENT_USED_FIELDS"
    return "CHANGED_USED_FIELDS"


def audit(mode: str, cache: Path) -> dict:
    evidence = json.loads((DOC / "DELIVERY_EVIDENCE_R8.json").read_text(encoding="utf-8"))
    claims = json.loads((DOC / "CLAIM_LEDGER_R8.json").read_text(encoding="utf-8"))["claims"]
    used_trip_ids = {str(c["value"]) for c in claims if c["metric"] in {"outbound_trip_id", "return_trip_id"}}
    used_stop_ids = {str(c["value"]) for c in claims if c["metric"] in {"outbound_stop_id", "return_stop_id"}}
    for case in ("MAIN", "VARIATION_TIME", "VARIATION_DURATION"):
        itinerary = evidence["cases"][case]["provider_result"]["itinerary"]
        used_stop_ids |= {itinerary["outbound"]["from_stop_id"], itinerary["outbound"]["to_stop_id"],
                          itinerary["return"]["from_stop_id"], itinerary["return"]["to_stop_id"]}
    health = json.loads((ROOT / "prototypes/ir_y_volver/snapshots/official-goierrialdea-go01-health-r5-20260929.json").read_text(encoding="utf-8"))
    node_ids = {str(node) for links in health["walking_links"].values() for link in links.values() for node in link["node_ids"]}
    way_ids = {str(way) for links in health["walking_links"].values() for link in links.values() for way in link["way_ids"]}
    retrieved = {source["source_id"]: source["retrieved_date"] for source in health["sources"] if source["source_id"] != "MODEL"}
    report = {"audit_version": "r9.0", "mode": mode, "checked_at_utc": datetime.now(timezone.utc).isoformat(),
              "runtime_changed": False, "automatic_pin_replacement": False, "sources": []}
    fetches = {}
    if mode == "online":
        cache.mkdir(parents=True, exist_ok=True)
        suffix = {"GTFS": ".zip", "HEALTH_REGISTRY": ".xlsx", "HEALTH_PAGE": ".html", "PADI_2026": ".pdf", "OSM": ".osm"}
        for source_id, url in URLS.items():
            target = cache / f"{source_id.lower()}{suffix[source_id]}"
            fetches[source_id] = fetch(url, target)
    pinned_gtfs = gtfs_signature(PINNED["GTFS"], used_trip_ids, used_stop_ids)
    pinned_health = health_record(PINNED["HEALTH_REGISTRY"])
    pinned_page = page_signature(PINNED["HEALTH_PAGE"])
    pinned_osm = osm_signature(PINNED["OSM"], node_ids, way_ids)
    for source_id in URLS:
        base = {"source_id": source_id, "url": URLS[source_id], "pinned_acquisition_timestamp": retrieved[source_id],
                "pinned_sha256": source_sha(PINNED[source_id], source_id), "pinned_local_artifact_sha256": source_sha(PINNED[source_id], source_id),
                "used_fields": [], "consequence": "No silent pin replacement.",
                "whether_runtime_rebuild_is_required": False, "whether_human_review_is_required": False}
        if mode == "offline":
            base.update({"classification": "INDETERMINATE", "semantic_comparison": "Online source not consulted in offline mode."})
        else:
            observed = fetches[source_id]
            base["online_observation"] = observed
            if observed["status"] != "FETCHED":
                base.update({"classification": "SOURCE_UNAVAILABLE", "semantic_comparison": "Fetch did not produce source bytes.",
                             "whether_human_review_is_required": True})
            else:
                current = cache / Path(observed["output"]).name
                base["current_sha256"] = observed["sha256"]
                if source_id == "GTFS":
                    now = gtfs_signature(current, used_trip_ids, used_stop_ids)
                    base.update({"used_fields": ["GO01", "used trip/stop IDs", "stop names", "2026-09-29 calendar", "arrival/departure", "pickup/dropoff", "timepoint"],
                                 "pinned_table_hashes": pinned_gtfs["canonical_table_sha256"], "current_table_hashes": now["canonical_table_sha256"],
                                 "pinned_used_fields_sha256": pinned_gtfs["used_fields_sha256"], "current_used_fields_sha256": now["used_fields_sha256"]})
                    metadata_only = pinned_gtfs["canonical_table_sha256"] == now["canonical_table_sha256"]
                    base["classification"] = classification(base["pinned_sha256"], observed["sha256"], pinned_gtfs["used_fields_sha256"], now["used_fields_sha256"], metadata_only=metadata_only)
                    base["semantic_comparison"] = "Canonical used-field signature compared; ZIP metadata ignored for table hashes."
                elif source_id == "HEALTH_REGISTRY":
                    now = health_record(current)
                    base.update({"used_fields": list(pinned_health), "pinned_used_fields": pinned_health, "current_used_fields": now,
                                 "classification": classification(base["pinned_sha256"], observed["sha256"], pinned_health, now),
                                 "semantic_comparison": "entityBC631DA5 identity/name/address/coordinates/telephone compared."})
                elif source_id == "HEALTH_PAGE":
                    now = page_signature(current)
                    base.update({"used_fields": list(pinned_page), "pinned_used_fields": pinned_page, "current_used_fields": now,
                                 "classification": classification(base["pinned_sha256"], observed["sha256"], pinned_page, now),
                                 "semantic_comparison": "Centre name/address/telephone presence compared; cosmetic HTML ignored."})
                elif source_id == "PADI_2026":
                    status = "UNCHANGED_BYTES" if base["pinned_sha256"] == observed["sha256"] else "INDETERMINATE"
                    base.update({"used_fields": ["Beasain PADI row/address/telephone"], "classification": status,
                                 "semantic_comparison": "Byte identity checked; changed PDF requires controlled row extraction/human review."})
                else:
                    now = osm_signature(current, node_ids, way_ids)
                    complete = len(now["used_nodes"]) == len(node_ids) and len(now["used_ways"]) == len(way_ids)
                    base.update({"used_fields": ["used node IDs", "used way IDs", "coordinates", "topology", "relevant tags"],
                                 "pinned_used_fields_sha256": pinned_osm["used_fields_sha256"], "current_used_fields_sha256": now["used_fields_sha256"],
                                 "classification": classification(base["pinned_sha256"], observed["sha256"], pinned_osm["used_fields_sha256"], now["used_fields_sha256"]),
                                 "semantic_comparison": f"Pinned route node/way topology compared; all used IDs present={complete}."})
                if base["classification"] in {"CHANGED_USED_FIELDS", "INDETERMINATE"}:
                    base["whether_human_review_is_required"] = True
                if base["classification"] == "CHANGED_USED_FIELDS":
                    base["consequence"] = "Review required; rebuild only after explicit decision. Historical evidence remains pinned."
        assert base["classification"] in DRIFT
        report["sources"].append(base)
    report["summary"] = {status: sum(row["classification"] == status for row in report["sources"]) for status in sorted(DRIFT)}
    report["human_review_required"] = any(row["whether_human_review_is_required"] for row in report["sources"])
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--offline", action="store_true")
    parser.add_argument("--online", action="store_true")
    parser.add_argument("--cache-dir", type=Path, default=ROOT / "work/r9-upstream")
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if args.offline == args.online:
        parser.error("select exactly one of --offline or --online")
    report = audit("online" if args.online else "offline", args.cache_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps({"mode": report["mode"], "summary": report["summary"], "human_review_required": report["human_review_required"]}))


if __name__ == "__main__":
    main()
